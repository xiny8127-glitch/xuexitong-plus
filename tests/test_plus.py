import unittest
from io import BytesIO
from unittest.mock import patch
from pathlib import Path
import tempfile

from xuexitong_mcp.homework import parse_requirements, safe_url
from xuexitong_mcp.attachments import check_download_url, extract, download, MAX_BYTES

HTML = '''<html><body><nav>导航</nav><main>
<div class="workDescription">提交 PDF 报告<br>命名为学号_姓名</div>
<div class="questionLi"><div class="mark_name">1. 求解 x²<img src="/formula.png" alt="x²"></div>
<ul class="options"><li>A. 1</li><li>B. 2</li></ul>
<div class="question-item">嵌套</div></div>
<a href="/files/report.docx">报告模板.docx</a>
<a href="javascript:alert(1)">bad.pdf</a><img src="">
<script>secret()</script></main></body></html>'''


class ParsingTests(unittest.TestCase):
    def test_rich_page(self):
        d = parse_requirements(HTML, 'https://mooc1.chaoxing.com/work')
        self.assertEqual(len(d['questions']), 1)
        self.assertEqual(d['questions'][0]['options'], ['A. 1', 'B. 2'])
        self.assertIn('学号_姓名', d['requirements'])
        self.assertEqual(len(d['attachments']), 1)
        self.assertEqual(d['images'][0]['url'], 'https://mooc1.chaoxing.com/formula.png')
        self.assertEqual(len(d['images']), 1)
        self.assertNotIn('secret', d['full_text'])
        self.assertNotIn('导航', d['full_text'])

    def test_unknown_and_iframe(self):
        d = parse_requirements('<body>请完成报告<iframe src="x"></iframe></body>')
        self.assertIn('请完成报告', d['full_text'])
        self.assertGreaterEqual(len(d['warnings']), 3)

    def test_access_block(self):
        with self.assertRaises(ValueError):
            parse_requirements('<p>没有此页面访问权限</p>')

    def test_urls(self):
        for u in ['http://mooc1.chaoxing.com/a', 'https://chaoxing.com.evil.com/a', 'https://127.0.0.1/a', 'https://u:p@chaoxing.com/a']:
            with self.assertRaises(ValueError):
                check_download_url(u)
        check_download_url('https://d0.cldisk.com/a')
        self.assertEqual(safe_url('javascript:alert(1)', 'https://chaoxing.com'), '')

    def test_text_and_docx(self):
        self.assertEqual(extract('作业要求'.encode('gb18030'), '要求.txt')['text'], '作业要求')
        from docx import Document
        doc = Document()
        doc.add_paragraph('提交实验报告')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '截止周五'
        buf = BytesIO()
        doc.save(buf)
        result = extract(buf.getvalue(), '模板.docx')
        self.assertIn('截止周五', result['text'])

    def test_download_bounds_and_redirect(self):
        from unittest.mock import MagicMock
        session = MagicMock()
        response = session.get.return_value.__enter__.return_value
        response.is_redirect = True
        response.headers = {'Location': 'https://evil.example/a.pdf'}
        with self.assertRaises(ValueError):
            download(session, 'https://mooc1.chaoxing.com/download')
        self.assertEqual(session.get.call_count, 1)
        response.is_redirect = False
        response.headers = {'Content-Length': str(MAX_BYTES+1)}
        with self.assertRaises(ValueError):
            download(session, 'https://mooc1.chaoxing.com/download')

    def test_ambiguous_course(self):
        from xuexitong_mcp.client import ChaoxingClient
        client = ChaoxingClient.__new__(ChaoxingClient)
        courses = [{'name': '数学', 'courseid': '1', 'clazzid': '2'}, {'name': '数学', 'courseid': '1', 'clazzid': '3'}]
        with patch.object(client, 'fetch_courses', return_value=courses):
            with self.assertRaises(ValueError):
                client.find_course('数学')
            self.assertEqual(client.find_course('1:3')['clazzid'], '3')

    def test_pdf_without_text(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        buf = BytesIO()
        writer.write(buf)
        result = extract(buf.getvalue(), '扫描件.pdf')
        self.assertEqual(result['text'], '')
        self.assertTrue(result['warnings'])

    def test_mcp_tools_and_export(self):
        from xuexitong_mcp import plus
        d = parse_requirements(HTML, 'https://mooc1.chaoxing.com/work')
        d.update(title='实验一', course='程序设计', status='未交', deadline=None)
        with patch.object(plus, 'get_data', side_effect=lambda *a: dict(d)):
            page = plus.read_homework_requirements('程序设计', 1, question_limit=1)
            self.assertEqual(page['total_questions_parsed'], 1)
            with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as temp:
                result = plus.export_homework_requirements('程序设计', 1, temp)
                self.assertIn('提交 PDF', Path(result['markdown']).read_text(encoding='utf-8'))
                self.assertTrue(Path(result['json']).exists())


if __name__ == '__main__':
    unittest.main()
