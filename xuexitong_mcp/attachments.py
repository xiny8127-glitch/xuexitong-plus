"""Bounded platform attachment downloads and text extraction."""
from io import BytesIO
from urllib.parse import urlsplit, urljoin
import zipfile

MAX_BYTES = 20 * 1024 * 1024


def check_download_url(url):
    p = urlsplit(url)
    host = (p.hostname or '').lower()
    if p.scheme != 'https' or p.username or p.port not in (None, 443) or not any(host == d or host.endswith('.'+d) for d in ('chaoxing.com', 'cldisk.com')):
        raise ValueError('附件仅允许 HTTPS 学习通/超星云盘域名')


def download(session, url):
    for _ in range(6):
        check_download_url(url)
        with session.get(url, stream=True, allow_redirects=False, timeout=(10, 60), headers={'Referer': 'https://pan-yz.chaoxing.com/'}) as r:
            if r.is_redirect:
                url = urljoin(url, r.headers['Location'])
                continue
            r.raise_for_status()
            if int(r.headers.get('Content-Length', 0)) > MAX_BYTES:
                raise ValueError('附件超过 20MB')
            buf = bytearray()
            for chunk in r.iter_content(65536):
                buf.extend(chunk)
                if len(buf) > MAX_BYTES:
                    raise ValueError('附件超过 20MB')
            if 'text/html' in r.headers.get('Content-Type', '').lower():
                raise ValueError('返回预览或登录网页，需在浏览器中查看附件')
            return bytes(buf)
    raise ValueError('附件重定向过多')


def extract(data: bytes, name: str) -> dict:
    warnings = []
    if data.startswith(b'%PDF'):
        from pypdf import PdfReader
        pages = PdfReader(BytesIO(data)).pages
        text = '\n\n'.join((p.extract_text() or '') for p in pages[:100])
        if len(pages) > 100:
            warnings.append('仅提取前 100 页')
        if not text.strip():
            warnings.append('PDF 无文字层，可能需要 OCR')
    elif data.startswith(b'PK'):
        with zipfile.ZipFile(BytesIO(data)) as z:
            if 'word/document.xml' not in z.namelist():
                raise ValueError('该压缩文件不是 DOCX，不进行解压')
            if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                raise ValueError('DOCX 解压大小超过限制')
        from docx import Document
        doc = Document(BytesIO(data))
        text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        warnings.append('仅提取段落和表格文字，未解析图片、公式、批注')
    elif name.lower().endswith(('.txt', '.csv', '.md')):
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('gb18030', errors='replace')
    else:
        raise ValueError('支持 PDF/DOCX/TXT/CSV/MD；其他格式请下载后查看')
    return {'text': text[:100000], 'truncated': len(text)>100000, 'warnings': warnings}
